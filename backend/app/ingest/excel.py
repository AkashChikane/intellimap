from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from .headers import EXPECTED_SHEETS, canonical_sheet, map_headers


def _cell(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def parse_workbook(path: str | Path) -> dict:
    wb = load_workbook(path, data_only=True, read_only=True)
    found = {}
    unknown_sheets = []
    for name in wb.sheetnames:
        canonical = canonical_sheet(name)
        if not canonical:
            unknown_sheets.append(name)
            continue
        found.setdefault(canonical, name)

    sheets = []
    for expected in EXPECTED_SHEETS:
        actual = found.get(expected)
        if not actual:
            sheets.append(
                {
                    "sheet": expected,
                    "workbook_name": None,
                    "status": "missing",
                    "row_count": 0,
                    "headers": [],
                    "mapped_headers": {},
                    "unmapped_headers": [],
                    "missing_fields": [],
                    "rows": [],
                }
            )
            continue
        ws = wb[actual]
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration:
            header_row = []
        headers = [_cell(h) for h in header_row]
        mapping = map_headers(expected, headers)
        records = []
        excel_row = 1
        for raw in rows_iter:
            excel_row += 1
            if raw is None or all(v is None or str(v).strip() == "" for v in raw):
                continue
            payload = {}
            original = {}
            for idx, header in enumerate(headers):
                val = _cell(raw[idx]) if idx < len(raw) else ""
                original[header] = val
                field = mapping["mapped"].get(header)
                if field:
                    payload[field] = val
            records.append(
                {
                    "source_row": excel_row,
                    "payload": payload,
                    "original": original,
                }
            )
        status = "ok"
        if mapping["missing"]:
            status = "partial"
        sheets.append(
            {
                "sheet": expected,
                "workbook_name": actual,
                "status": status,
                "row_count": len(records),
                "headers": headers,
                "mapped_headers": mapping["mapped"],
                "unmapped_headers": mapping["unmapped"],
                "missing_fields": mapping["missing"],
                "rows": records,
            }
        )
    wb.close()
    return {
        "filename": Path(path).name,
        "sheets": sheets,
        "unknown_sheets": unknown_sheets,
    }
