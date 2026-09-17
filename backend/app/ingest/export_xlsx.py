from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from ..db.store import Store
from .headers import EXPECTED_SHEETS, pretty_header
from .pipeline import TABLE_KEYS

HEADER_FILL = PatternFill("solid", fgColor="1F2F57")
HEADER_FONT = Font(color="FDFAF9", bold=True, name="Calibri")
CELL_FONT = Font(name="Calibri", color="1F2F57")
OMIT_FONT = Font(name="Calibri", color="A8A8A8", italic=True)


def workbook_bytes(store: Store, run_id: str) -> bytes:
    omitted = set(store.omitted_sheets(run_id))
    wb = Workbook()
    wb.remove(wb.active)
    for sheet in EXPECTED_SHEETS:
        if sheet in omitted:
            continue
        if sheet not in TABLE_KEYS:
            continue
        rows = store.raw_rows(run_id, sheet)
        headers = _headers(sheet, rows)
        ws = wb.create_sheet(sheet)
        for col, header in enumerate(headers, 1):
            cell = ws.cell(1, col, header)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="left")
        for r_idx, row in enumerate(rows, 2):
            original = row.get("original") or {}
            payload = row.get("payload") or {}
            for c_idx, header in enumerate(headers, 1):
                value = original.get(header)
                if value is None:
                    field = _field_for_header(sheet, header)
                    value = payload.get(field, "") if field else ""
                cell = ws.cell(r_idx, c_idx, value)
                cell.font = CELL_FONT
        for col, header in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(col)].width = min(max(len(header) + 4, 14), 36)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    if not wb.sheetnames:
        ws = wb.create_sheet("Empty")
        ws["A1"] = "All sheets were omitted in review."
        ws["A1"].font = OMIT_FONT
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _headers(sheet: str, rows: list[dict]) -> list[str]:
    seen = []
    for row in rows:
        for header in (row.get("original") or {}):
            if header and header not in seen:
                seen.append(header)
    if seen:
        return seen
    from .headers import FIELD_ALIASES

    fields = []
    for field in (FIELD_ALIASES.get(sheet) or {}).values():
        if field not in fields:
            fields.append(field)
    return [pretty_header(f) for f in fields]


def _field_for_header(sheet: str, header: str) -> str | None:
    from .headers import FIELD_ALIASES, normalize_token

    aliases = FIELD_ALIASES.get(sheet) or {}
    return aliases.get(normalize_token(header))
