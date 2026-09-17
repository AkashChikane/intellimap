"""Build the bundled synthetic architecture workbook.

IDs in this file are sample DATA, never referenced by validation or graph code.
"""
from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "data" / "IntelliMap_Architecture_Landscape.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="1F2F57")
HEADER_FONT = Font(color="FDFAF9", bold=True, name="Calibri")
CELL_FONT = Font(name="Calibri", color="1F2F57")


def _app(
    aid,
    name,
    desc,
    domain,
    crit,
    status,
    start,
    end,
    hosting,
    vendor,
    owner,
    cc,
):
    return {
        "ApplicationID": aid,
        "ApplicationName": name,
        "Description": desc,
        "BusinessDomain": domain,
        "BusinessCriticality": crit,
        "Lifecyclestatus": status,
        "LifecycleStartDate": start,
        "LifecycleEndDate": end,
        "Hosting": hosting,
        "Vendor Type": vendor,
        "OwnerEmployeeID": owner,
        "CostCenter": cc,
    }


def applications():
    rows = [
        _app("APP-0001", "Customer Portal", "External order capture and status for dealers and retail customers.", "Order-to-Cash", "High", "Active", "2019-03-01", "", "SaaS", "Commercial", "E-1041", "CC-OTC"),
        _app("APP-0002", "Configure Price Quote", "Vehicle configuration, discounting, and quote generation.", "Order-to-Cash", "High", "Active", "2020-01-15", "", "Cloud", "Commercial", "E-1041", "CC-OTC"),
        _app("APP-0003", "Product Catalog", "Sellable vehicle and option master used by CPQ and portal.", "Order-to-Cash", "Medium", "Active", "2018-06-01", "", "On-prem", "Custom", "E-1102", "CC-OTC"),
        _app("APP-0004", "Pricing Engine", "List, option, and campaign pricing services.", "Order-to-Cash", "High", "Active", "2021-02-01", "", "Cloud", "Custom", "E-1102", "CC-OTC"),
        _app("APP-0005", "Order Management System", "System of record for vehicle orders, holds, and fulfilment triggers.", "Order-to-Cash", "Mission Critical", "Active", "2016-09-01", "", "On-prem", "Custom", "E-1007", "CC-OTC"),
        _app("APP-0006", "Billing System", "Customer invoicing from fulfilled orders.", "Order-to-Cash", "High", "Active", "2017-04-01", "", "On-prem", "Commercial", "E-1220", "CC-FIN"),
        _app("APP-0007", "Invoice Management", "Invoice presentation, disputes, and e-invoicing.", "Order-to-Cash", "Medium", "Active", "2022-01-10", "", "SaaS", "Commercial", "E-1220", "CC-FIN"),
        _app("APP-0008", "Collections Workbench", "Dunning and collections case management.", "Order-to-Cash", "Medium", "Active", "2019-11-01", "", "Cloud", "Commercial", "E-1221", "CC-FIN"),
        _app("APP-0020", "Manufacturing Execution", "Shop-floor execution, genealogy, and work-order confirmations.", "Manufacturing", "Mission Critical", "Active", "2015-01-01", "", "On-prem", "Commercial", "E-2001", "CC-MFG"),
        _app("APP-0021", "Product Lifecycle Management", "BOM, change control, and engineering documentation.", "Manufacturing", "High", "Active", "2014-08-01", "", "On-prem", "Commercial", "E-2008", "CC-ENG"),
        _app("APP-0022", "Quality Management System", "NCs, CAPA, and incoming inspection.", "Manufacturing", "High", "Active", "2018-02-01", "", "Cloud", "Commercial", "E-2014", "CC-MFG"),
        _app("APP-0023", "Shop Floor Control", "Line sequencing and station instructions.", "Manufacturing", "High", "Active", "2016-05-01", "", "On-prem", "Custom", "E-2001", "CC-MFG"),
        _app("APP-0024", "ERP Core", "Finance, inventory valuation, and manufacturing orders.", "Manufacturing", "Mission Critical", "Active", "2012-01-01", "", "On-prem", "Commercial", "E-3001", "CC-ERP"),
        _app("APP-0025", "Production Scheduling", "Finite scheduling across plants.", "Manufacturing", "High", "Active", "2020-07-01", "", "Cloud", "Commercial", "E-2003", "CC-MFG"),
        _app("APP-0026", "Maintenance Management", "Preventive and breakdown maintenance.", "Manufacturing", "Medium", "Active", "2017-09-01", "", "SaaS", "Commercial", "E-2050", "CC-MFG"),
        _app("APP-0030", "Inventory Management", "Plant and depot stock positions.", "Supply Chain", "High", "Active", "2013-04-01", "", "On-prem", "Commercial", "E-3104", "CC-SCM"),
        _app("APP-0031", "Warehouse Management", "DC operations, waves, and shipping.", "Supply Chain", "High", "Active", "2016-02-01", "", "On-prem", "Commercial", "E-3104", "CC-SCM"),
        _app("APP-0032", "Transportation Management", "Carrier booking and freight audit.", "Supply Chain", "Medium", "Active", "2019-08-01", "", "SaaS", "Commercial", "E-3110", "CC-SCM"),
        _app("APP-0033", "Supplier Portal", "ASN, scheduling agreements, and supplier quality.", "Supply Chain", "Medium", "Active", "2021-03-01", "", "SaaS", "Commercial", "E-3122", "CC-SCM"),
        _app("APP-0034", "Procurement Suite", "P2P, contracts, and catalogs.", "Supply Chain", "High", "Active", "2018-01-01", "", "SaaS", "Commercial", "E-3120", "CC-SCM"),
        _app("APP-0035", "Demand Planning", "Statistical and consensus demand.", "Supply Chain", "Medium", "Active", "2022-05-01", "", "Cloud", "Commercial", "E-3130", "CC-SCM"),
        _app("APP-0036", "Yard Management", "Gate, dock, and trailer yard control.", "Supply Chain", "Low", "Active", "2023-01-01", "", "SaaS", "Commercial", "E-3109", "CC-SCM"),
        _app("APP-0040", "Legacy EDI Gateway", "Historic VAN/EDI translator still used for a few OEM partners.", "Finance", "High", "Retired", "2008-01-01", "2024-12-31", "On-prem", "Custom", "E-4099", "CC-INT"),
        _app("APP-0041", "General Ledger", "Corporate ledger and close.", "Finance", "Mission Critical", "Active", "2012-01-01", "", "On-prem", "Commercial", "E-4001", "CC-FIN"),
        _app("APP-0042", "Accounts Payable", "Vendor invoices and payments.", "Finance", "High", "Active", "2012-01-01", "", "On-prem", "Commercial", "E-4002", "CC-FIN"),
        _app("APP-0043", "Accounts Receivable", "Customer open items and cash application.", "Finance", "High", "Active", "2012-01-01", "", "On-prem", "Commercial", "E-4003", "CC-FIN"),
        _app("APP-0044", "Treasury", "Cash, payments, and bank connectivity.", "Finance", "High", "Active", "2019-06-01", "", "SaaS", "Commercial", "E-4010", "CC-FIN"),
        _app("APP-0045", "Tax Engine", "Indirect tax determination.", "Finance", "Medium", "Active", "2020-10-01", "", "SaaS", "Commercial", "E-4011", "CC-FIN"),
        _app("APP-0050", "HRIS Core", "Employee master, org, and workforce admin.", "HR", "High", "Active", "2018-03-01", "", "SaaS", "Commercial", "E-5001", "CC-HR"),
        _app("APP-0051", "Payroll", "Payroll calculation and disbursement.", "HR", "Mission Critical", "Active", "2018-03-01", "", "SaaS", "Commercial", "E-5002", "CC-HR"),
        _app("APP-0052", "Recruiting", "Requisitions and candidate pipeline.", "HR", "Medium", "Active", "2026-01-01", "2024-01-01", "SaaS", "Commercial", "E-5003", "CC-HR"),
        _app("APP-0053", "Time and Attendance", "Clockings and absence.", "HR", "Medium", "Active", "2019-01-01", "", "SaaS", "Commercial", "E-5004", "CC-HR"),
        _app("APP-0054", "Learning Management", "Training catalog and completions. No current integrations.", "HR", "Low", "Active", "2021-09-01", "", "SaaS", "Commercial", "", "CC-HR"),
        _app("APP-0060", "CRM", "Accounts, opportunities, and dealer relationships.", "Customer", "High", "Active", "2017-05-01", "", "SaaS", "Commercial", "E-6001", "CC-CX"),
        _app("APP-0061", "Marketing Automation", "Campaigns and lead scoring.", "Customer", "Medium", "Active", "2019-02-01", "2023-06-01", "SaaS", "Commercial", "", "CC-CX"),
        _app("APP-0062", "Customer Master Hub", "Golden customer and party records.", "Customer", "High", "Active", "2020-04-01", "", "Cloud", "Custom", "E-6002", "CC-CX"),
        _app("APP-0063", "Field Service", "Technician dispatch and installed base.", "Customer", "Medium", "Active", "2021-11-01", "", "SaaS", "Commercial", "E-6008", "CC-CX"),
        _app("APP-0064", "Loyalty Platform", "Points, tiers, and rewards.", "Customer", "Low", "Active", "2022-08-01", "", "SaaS", "Commercial", "E-6011", "CC-CX"),
        _app("APP-0065", "Contact Center", "Omnichannel agent desktop.", "Customer", "Medium", "Active", "2018-12-01", "", "Cloud", "Commercial", "E-6012", "CC-CX"),
        _app("APP-0070", "Enterprise Data Warehouse", "Curated enterprise reporting marts.", "Data Platform", "High", "Active", "2014-01-01", "", "On-prem", "Custom", "E-7001", "CC-DATA"),
        _app("APP-0071", "Data Lake", "Raw and refined object store.", "Data Platform", "High", "Active", "2020-01-01", "", "Cloud", "Commercial", "E-7001", "CC-DATA"),
        _app("APP-0072", "Master Data Management", "Product, location, and party mastering.", "Data Platform", "High", "Active", "2019-07-01", "", "Cloud", "Commercial", "E-7004", "CC-DATA"),
        _app("APP-0073", "Analytics Workbench", "Self-service BI and notebooks.", "Data Platform", "Medium", "Active", "2021-04-01", "", "SaaS", "Commercial", "E-7008", "CC-DATA"),
        _app("APP-0074", "Integration Streaming Bus", "Enterprise event backbone.", "Data Platform", "High", "Active", "2022-02-01", "", "Cloud", "Custom", "E-7010", "CC-INT"),
        _app("APP-0075", "Data Catalog", "Dataset inventory and lineage.", "Data Platform", "Low", "Active", "2023-03-01", "", "SaaS", "Commercial", "E-7009", "CC-DATA"),
        _app("APP-0080", "API Gateway", "Managed external and internal APIs.", "Integration", "High", "Active", "2020-06-01", "", "Cloud", "Commercial", "E-8001", "CC-INT"),
        _app("APP-0081", "iPaaS", "Low-code integration platform.", "Integration", "High", "Active", "2021-01-01", "", "SaaS", "Commercial", "E-8001", "CC-INT"),
        _app("APP-0082", "File Transfer Gateway", "Managed file transfer.", "Integration", "Medium", "Active", "2015-01-01", "", "On-prem", "Commercial", "E-8004", "CC-INT"),
        _app("APP-0083", "Event Broker", "Pub/sub topics for domain events.", "Integration", "High", "Active", "2022-02-01", "", "Cloud", "Commercial", "E-7010", "CC-INT"),
        _app("APP-0084", "Partner B2B Hub", "Partner onboarding and B2B traffic.", "Integration", "Medium", "Active", "2016-08-01", "", "On-prem", "Commercial", "E-8006", "CC-INT"),
        _app("APP-0090", "Identity Provider", "Workforce and partner SSO.", "Identity/Security", "Mission Critical", "Active", "2019-01-01", "", "SaaS", "Commercial", "E-9001", "CC-SEC"),
        _app("APP-0091", "Privileged Access", "Vaulted admin sessions.", "Identity/Security", "High", "Active", "2021-06-01", "", "SaaS", "Commercial", "E-9002", "CC-SEC"),
        _app("APP-0092", "SIEM", "Security event correlation.", "Identity/Security", "High", "Active", "2018-04-01", "", "Cloud", "Commercial", "E-9003", "CC-SEC"),
        _app("APP-0093", "Secrets Vault", "Application secrets and certificates.", "Identity/Security", "High", "Active", "2020-09-01", "", "Cloud", "Commercial", "E-9002", "CC-SEC"),
        _app("APP-0094", "GRC Platform", "Controls, issues, and policy attestations.", "Identity/Security", "Medium", "Active", "2023-05-01", "", "SaaS", "Commercial", "E-9008", "CC-SEC"),
        _app("APP-0095", "Dealer DMS Adapter", "Adapter used by selected national dealer systems.", "Order-to-Cash", "Medium", "Active", "2022-09-01", "", "Cloud", "Custom", "E-1044", "CC-OTC"),
        _app("APP-0096", "Plant Historian", "Time-series telemetry from production lines.", "Manufacturing", "Medium", "Active", "2017-01-01", "", "On-prem", "Commercial", "E-2060", "CC-MFG"),
        _app("APP-0097", "Carbon Accounting", "Emissions factors and sustainability reporting.", "Finance", "Low", "Active", "2024-02-01", "", "SaaS", "Commercial", "E-4020", "CC-FIN"),
        _app("APP-0098", "Expense Management", "T&E capture and policy checks.", "Finance", "Low", "Active", "2020-03-01", "", "SaaS", "Commercial", "E-4022", "CC-FIN"),
        _app("APP-0099", "Contract Lifecycle", "Legal agreements and clause library.", "Supply Chain", "Medium", "Active", "2021-12-01", "", "SaaS", "Commercial", "E-3128", "CC-SCM"),
    ]
    return rows


def relationships():
    def r(rid, src, sname, rtype, tgt, tname, direction, crit):
        return {
            "RelationshipID": rid,
            "SourceApplicationID": src,
            "SourceApplicationName": sname,
            "RelationshipType": rtype,
            "TargetApplicationID": tgt,
            "TargetApplicationName": tname,
            "Direction": direction,
            "DependencyCriticality": crit,
        }

    rows = [
        r("REL-0001", "APP-0001", "Customer Portal", "depends_on", "APP-0005", "Order Management System", "source_to_target", "High"),
        r("REL-0002", "APP-0001", "Customer Portal", "depends_on", "APP-0002", "Configure Price Quote", "source_to_target", "High"),
        r("REL-0003", "APP-0002", "Configure Price Quote", "depends_on", "APP-0003", "Product Catalog", "source_to_target", "High"),
        r("REL-0004", "APP-0002", "Configure Price Quote", "depends_on", "APP-0004", "Pricing Engine", "source_to_target", "High"),
        r("REL-0005", "APP-0002", "Configure Price Quote", "depends_on", "APP-0005", "Order Management System", "source_to_target", "Medium"),
        r("REL-0006", "APP-0005", "Order Management System", "depends_on", "APP-0024", "ERP Core", "source_to_target", "Mission Critical"),
        r("REL-0007", "APP-0005", "Order Management System", "depends_on", "APP-0030", "Inventory Management", "source_to_target", "High"),
        r("REL-0008", "APP-0005", "Order Management System", "depends_on", "APP-0004", "Pricing Engine", "source_to_target", "Medium"),
        r("REL-0009", "APP-0005", "OMS", "depends_on", "APP-0080", "API Gateway", "source_to_target", "High"),
        r("REL-0010", "APP-0006", "Billing System", "depends_on", "APP-0005", "Order Management System", "source_to_target", "High"),
        r("REL-0011", "APP-0006", "Billing System", "depends_on", "APP-0043", "Accounts Receivable", "source_to_target", "High"),
        r("REL-0012", "APP-0007", "Invoice Management", "depends_on", "APP-0006", "Billing System", "source_to_target", "High"),
        r("REL-0013", "APP-0008", "Collections Workbench", "depends_on", "APP-0043", "Accounts Receivable", "source_to_target", "High"),
        r("REL-0014", "APP-0020", "Manufacturing Execution", "depends_on", "APP-0024", "ERP Core", "source_to_target", "Mission Critical"),
        r("REL-0015", "APP-0020", "Manufacturing Execution", "depends_on", "APP-0023", "Shop Floor Control", "source_to_target", "High"),
        r("REL-0016", "APP-0023", "Shop Floor Control", "depends_on", "APP-0025", "Production Scheduling", "source_to_target", "High"),
        r("REL-0017", "APP-0025", "Production Scheduling", "depends_on", "APP-0024", "ERP Core", "source_to_target", "High"),
        r("REL-0018", "APP-0021", "Product Lifecycle Management", "feeds", "APP-0003", "Product Catalog", "source_to_target", "High"),
        r("REL-0019", "APP-0021", "Product Lifecycle Management", "feeds", "APP-0024", "ERP Core", "source_to_target", "High"),
        r("REL-0020", "APP-0022", "Quality Management System", "depends_on", "APP-0020", "Manufacturing Execution", "source_to_target", "Medium"),
        r("REL-0021", "APP-0026", "Maintenance Management", "depends_on", "APP-0020", "Manufacturing Execution", "source_to_target", "Low"),
        r("REL-0022", "APP-0030", "Inventory Management", "depends_on", "APP-0024", "ERP Core", "source_to_target", "High"),
        r("REL-0023", "APP-0030", "Inventory Management", "depends_on", "APP-0031", "Warehouse Management", "source_to_target", "High"),
        r("REL-0024", "APP-0031", "Warehouse Management", "depends_on", "APP-0005", "Order Management System", "source_to_target", "Medium"),
        r("REL-0025", "APP-0031", "Warehouse Management", "depends_on", "APP-0032", "Transportation Management", "source_to_target", "Medium"),
        r("REL-0026", "APP-0032", "Transportation Management", "depends_on", "APP-0084", "Partner B2B Hub", "source_to_target", "Low"),
        r("REL-0027", "APP-0034", "Procurement Suite", "depends_on", "APP-0024", "ERP Core", "source_to_target", "High"),
        r("REL-0028", "APP-0034", "Procurement Suite", "depends_on", "APP-0033", "Supplier Portal", "source_to_target", "Medium"),
        r("REL-0029", "APP-0035", "Demand Planning", "depends_on", "APP-0070", "Enterprise Data Warehouse", "source_to_target", "Medium"),
        r("REL-0030", "APP-0035", "Demand Planning", "feeds", "APP-0025", "Production Scheduling", "source_to_target", "Medium"),
        r("REL-0031", "APP-0040", "Legacy EDI Gateway", "feeds", "APP-0005", "Order Management System", "source_to_target", "High"),
        r("REL-0032", "APP-0005", "Order Management System", "depends_on", "APP-0040", "Legacy EDI Gateway", "source_to_target", "High"),
        r("REL-0033", "APP-0041", "General Ledger", "depends_on", "APP-0024", "ERP Core", "source_to_target", "Mission Critical"),
        r("REL-0034", "APP-0042", "Accounts Payable", "depends_on", "APP-0024", "ERP Core", "source_to_target", "High"),
        r("REL-0035", "APP-0043", "Accounts Receivable", "depends_on", "APP-0024", "ERP Core", "source_to_target", "High"),
        r("REL-0036", "APP-0044", "Treasury", "depends_on", "APP-0041", "General Ledger", "source_to_target", "High"),
        r("REL-0037", "APP-0045", "Tax Engine", "used_by", "APP-0006", "Billing System", "target_to_source", "Medium"),
        r("REL-0038", "APP-0051", "Payroll", "depends_on", "APP-0050", "HRIS Core", "source_to_target", "Mission Critical"),
        r("REL-0039", "APP-0053", "Time and Attendance", "feeds", "APP-0051", "Payroll", "source_to_target", "High"),
        r("REL-0040", "APP-0052", "Recruiting", "feeds", "APP-0050", "HRIS Core", "source_to_target", "Medium"),
        r("REL-0041", "APP-0060", "CRM", "feeds", "APP-0005", "Order Management System", "source_to_target", "High"),
        r("REL-0042", "APP-0060", "CRM", "depends_on", "APP-0062", "Customer Master Hub", "source_to_target", "High"),
        r("REL-0043", "APP-0061", "Marketing Automation", "feeds", "APP-0060", "CRM", "source_to_target", "Medium"),
        r("REL-0044", "APP-0065", "Contact Center", "depends_on", "APP-0060", "CRM", "source_to_target", "Medium"),
        r("REL-0045", "APP-0063", "Field Service", "depends_on", "APP-0062", "Customer Master Hub", "source_to_target", "Medium"),
        r("REL-0046", "APP-0063", "Field Service", "depends_on", "APP-0021", "Product Lifecycle Management", "source_to_target", "Low"),
        r("REL-0047", "APP-0064", "Loyalty Platform", "depends_on", "APP-0062", "Customer Master Hub", "source_to_target", "Low"),
        r("REL-0048", "APP-0070", "Enterprise Data Warehouse", "depends_on", "APP-0024", "ERP Core", "source_to_target", "High"),
        r("REL-0049", "APP-0070", "Enterprise Data Warehouse", "depends_on", "APP-0005", "Order Management System", "source_to_target", "High"),
        r("REL-0050", "APP-0071", "Data Lake", "depends_on", "APP-0074", "Integration Streaming Bus", "source_to_target", "High"),
        r("REL-0051", "APP-0073", "Analytics Workbench", "depends_on", "APP-0070", "Enterprise Data Warehouse", "source_to_target", "Medium"),
        r("REL-0052", "APP-0073", "Analytics Workbench", "depends_on", "APP-0071", "Data Lake", "source_to_target", "Medium"),
        r("REL-0053", "APP-0072", "Master Data Management", "feeds", "APP-0062", "Customer Master Hub", "source_to_target", "High"),
        r("REL-0054", "APP-0072", "Master Data Management", "feeds", "APP-0003", "Product Catalog", "source_to_target", "High"),
        r("REL-0055", "APP-0074", "Integration Streaming Bus", "depends_on", "APP-0083", "Event Broker", "source_to_target", "High"),
        r("REL-0056", "APP-0080", "API Gateway", "depends_on", "APP-0090", "Identity Provider", "source_to_target", "Mission Critical"),
        r("REL-0057", "APP-0081", "iPaaS", "depends_on", "APP-0080", "API Gateway", "source_to_target", "Medium"),
        r("REL-0058", "APP-0084", "Partner B2B Hub", "depends_on", "APP-0040", "Legacy EDI Gateway", "source_to_target", "High"),
        r("REL-0059", "APP-0084", "Partner B2B Hub", "depends_on", "APP-0082", "File Transfer Gateway", "source_to_target", "Medium"),
        r("REL-0060", "APP-0091", "Privileged Access", "depends_on", "APP-0090", "Identity Provider", "source_to_target", "High"),
        r("REL-0061", "APP-0092", "SIEM", "receives_from", "APP-0080", "API Gateway", "source_to_target", "Medium"),
        r("REL-0062", "APP-0093", "Secrets Vault", "used_by", "APP-0081", "iPaaS", "bidirectional", "High"),
        r("REL-0063", "APP-0005", "Order Management System", "depends_on", "APP-8824", "ERP Core", "source_to_target", "Mission Critical"),
        r("REL-0064", "APP-0095", "Dealer DMS Adapter", "feeds", "APP-0005", "Order Management System", "source_to_target", "Medium"),
        r("REL-0065", "APP-0096", "Plant Historian", "feeds", "APP-0071", "Data Lake", "source_to_target", "Low"),
        r("REL-0066", "APP-0097", "Carbon Accounting", "depends_on", "APP-0070", "Enterprise Data Warehouse", "source_to_target", "Low"),
        r("REL-0067", "APP-0098", "Expense Management", "feeds", "APP-0042", "Accounts Payable", "source_to_target", "Low"),
        r("REL-0068", "APP-0099", "Contract Lifecycle", "feeds", "APP-0034", "Procurement Suite", "source_to_target", "Medium"),
        r("REL-0069", "APP-0005", "Order Management System", "depends_on", "APP-0074", "Integration Streaming Bus", "source_to_target", "Medium"),
        r("REL-0070", "APP-0062", "Customer Master Hub", "depends_on", "APP-0074", "Integration Streaming Bus", "source_to_target", "Medium"),
        r("REL-0071", "APP-0024", "ERP Core", "depends_on", "APP-0090", "Identity Provider", "source_to_target", "High"),
        r("REL-0072", "APP-0005", "Order Management System", "depends_on", "APP-0090", "Identity Provider", "source_to_target", "High"),
        r("REL-0073", "APP-0075", "Data Catalog", "depends_on", "APP-0071", "Data Lake", "source_to_target", "Low"),
        r("REL-0074", "APP-0036", "Yard Management", "depends_on", "APP-0031", "Warehouse Management", "source_to_target", "Low"),
    ]
    return rows


def interfaces():
    def i(iid, name, prov, pname, cons, cname, proto, fmt, freq, status):
        return {
            "InterfaceID": iid,
            "InterfaceName": name,
            "ProviderApplicationID": prov,
            "ProviderApplicationName": pname,
            "ConsumerApplicationID": cons,
            "ConsumerApplicationName": cname,
            "Protocol": proto,
            "DataFormat": fmt,
            "Frequency": freq,
            "InterfaceStatus": status,
        }

    return [
        i("IF-0001", "Order-to-ERP Sync", "APP-0024", "ERP Core", "APP-0005", "Order Management System", "REST/HTTPS", "JSON", "Near real-time", "Active"),
        i("IF-0002", "Portal Order Submit", "APP-0005", "Order Management System", "APP-0001", "Customer Portal", "REST/HTTPS", "JSON", "Real-time", "Active"),
        i("IF-0003", "CPQ Quote Push", "APP-0005", "Order Management System", "APP-0002", "Configure Price Quote", "REST/HTTPS", "JSON", "Real-time", "Active"),
        i("IF-0004", "Catalog Read", "APP-0003", "Product Catalog", "APP-0002", "Configure Price Quote", "REST/HTTPS", "JSON", "Hourly", "Active"),
        i("IF-0005", "Price Lookup", "APP-0004", "Pricing Engine", "APP-0002", "Configure Price Quote", "REST/HTTPS", "JSON", "Real-time", "Active"),
        i("IF-0006", "Inventory Availability", "APP-0030", "Inventory Management", "APP-0005", "Order Management System", "REST/HTTPS", "JSON", "Near real-time", "Active"),
        i("IF-0007", "Billing Event", "APP-0006", "Billing System", "APP-0005", "Order Management System", "Events", "JSON", "Near real-time", "Active"),
        i("IF-0008", "AR Open Item", "APP-0043", "Accounts Receivable", "APP-0006", "Billing System", "Batch", "CSV", "Nightly", "Active"),
        i("IF-0009", "MES Confirmation", "APP-0024", "ERP Core", "APP-0020", "Manufacturing Execution", "RFC", "IDoc", "Near real-time", "Active"),
        i("IF-0010", "BOM Publish", "APP-0024", "ERP Core", "APP-0021", "Product Lifecycle Management", "Batch", "XML", "Nightly", "Active"),
        i("IF-0011", "Schedule Release", "APP-0020", "Manufacturing Execution", "APP-0025", "Production Scheduling", "REST/HTTPS", "JSON", "Hourly", "Active"),
        i("IF-0012", "WMS Pick Wave", "APP-0031", "Warehouse Management", "APP-0005", "Order Management System", "REST/HTTPS", "JSON", "Near real-time", "Active"),
        i("IF-0013", "ASN Inbound", "APP-0031", "Warehouse Management", "APP-0033", "Supplier Portal", "AS2", "EDIFACT", "Event", "Active"),
        i("IF-0014", "PO Export", "APP-0024", "ERP Core", "APP-0034", "Procurement Suite", "Batch", "XML", "Hourly", "Active"),
        i("IF-0015", "GL Journal", "APP-0041", "General Ledger", "APP-0024", "ERP Core", "Batch", "IDoc", "Nightly", "Active"),
        i("IF-0016", "Payment File", "APP-0044", "Treasury", "APP-0006", "Billing System", "SFTP", "XML", "Daily", "Active"),
        i("IF-0017", "Tax Determine", "APP-0045", "Tax Engine", "APP-0006", "Billing System", "REST/HTTPS", "JSON", "Real-time", "Active"),
        i("IF-0018", "Employee Master", "APP-0051", "Payroll", "APP-0050", "HRIS Core", "Batch", "JSON", "Nightly", "Active"),
        i("IF-0019", "Time Export", "APP-0051", "Payroll", "APP-0053", "Time and Attendance", "Batch", "CSV", "Weekly", "Active"),
        i("IF-0020", "Customer Golden Record", "APP-0062", "Customer Master Hub", "APP-0060", "CRM", "REST/HTTPS", "JSON", "Near real-time", "Active"),
        i("IF-0021", "Lead Handoff", "APP-0060", "CRM", "APP-0061", "Marketing Automation", "REST/HTTPS", "JSON", "Hourly", "Active"),
        i("IF-0022", "OMS Customer Snapshot", "APP-0005", "Order Management System", "APP-0062", "Customer Master Hub", "Events", "JSON", "Near real-time", "Active"),
        i("IF-0023", "EDW Order Facts", "APP-0070", "Enterprise Data Warehouse", "APP-0005", "Order Management System", "Batch", "Parquet", "Nightly", "Active"),
        i("IF-0024", "Event Ingest", "APP-0074", "Integration Streaming Bus", "APP-0005", "Order Management System", "Kafka", "Avro", "Real-time", "Active"),
        i("IF-0025", "IdP SAML", "APP-0090", "Identity Provider", "APP-0005", "Order Management System", "SAML", "XML", "Real-time", "Active"),
        i("IF-0026", "IdP OIDC", "APP-0090", "Identity Provider", "APP-0080", "API Gateway", "OIDC", "JWT", "Real-time", "Active"),
        i("IF-0027", "Partner EDI In", "APP-0040", "Legacy EDI Gateway", "APP-0084", "Partner B2B Hub", "VAN", "X12", "Event", "Deprecated"),
        i("IF-0028", "Dealer Order In", "APP-0005", "Order Management System", "APP-0095", "Dealer DMS Adapter", "REST/HTTPS", "JSON", "Near real-time", "Active"),
        i("IF-0029", "Historian Telemetry", "APP-0071", "Data Lake", "APP-0096", "Plant Historian", "MQTT", "JSON", "Streaming", "Active"),
        i("IF-0030", "Expense to AP", "APP-0042", "Accounts Payable", "APP-0098", "Expense Management", "Batch", "CSV", "Daily", "Active"),
        i("IF-0031", "Contract to P2P", "APP-0034", "Procurement Suite", "APP-0099", "Contract Lifecycle", "REST/HTTPS", "JSON", "Hourly", "Active"),
        i("IF-0032", "Field Service Installed Base", "APP-0063", "Field Service", "APP-0021", "Product Lifecycle Management", "REST/HTTPS", "JSON", "Daily", "Active"),
        i("IF-0033", "Loyalty Profile", "APP-0064", "Loyalty Platform", "APP-0062", "Customer Master Hub", "REST/HTTPS", "JSON", "Hourly", "Active"),
        i("IF-0034", "SIEM Gateway Logs", "APP-0092", "SIEM", "APP-0080", "API Gateway", "Syslog", "JSON", "Streaming", "Active"),
        i("IF-0035", "iPaaS Secrets", "APP-0093", "Secrets Vault", "APP-0081", "iPaaS", "REST/HTTPS", "JSON", "On demand", "Active"),
        i("IF-0036", "Demand Signal", "APP-0025", "Production Scheduling", "APP-0035", "Demand Planning", "Batch", "CSV", "Nightly", "Active"),
        i("IF-0037", "Quality NC", "APP-0022", "Quality Management System", "APP-0020", "Manufacturing Execution", "REST/HTTPS", "JSON", "Event", "Active"),
        i("IF-0038", "Yard Check-in", "APP-0036", "Yard Management", "APP-0031", "Warehouse Management", "REST/HTTPS", "JSON", "Event", "Active"),
        i("IF-0088", "Orphan Catalog Ping", "APP-0003", "Product Catalog", "APP-0075", "Data Catalog", "REST/HTTPS", "JSON", "Weekly", "Active"),
        i("IF-0040", "Contact Center Screen Pop", "APP-0065", "Contact Center", "APP-0060", "CRM", "REST/HTTPS", "JSON", "Real-time", "Active"),
    ]


def information_objects():
    def f(fid, obj, clas, src, sname, tgt, tname, op, iid):
        return {
            "FlowID": fid,
            "InformationObject": obj,
            "Classification": clas,
            "SourceApplicationID": src,
            "SourceApplicationName": sname,
            "TargetApplicationID": tgt,
            "TargetApplicationName": tname,
            "Operation": op,
            "InterfaceID": iid,
        }

    return [
        f("FLOW-0001", "Customer Quote", "Internal", "APP-0002", "Configure Price Quote", "APP-0005", "Order Management System", "Create", "IF-0003"),
        f("FLOW-0002", "Web Order", "Internal", "APP-0001", "Customer Portal", "APP-0005", "Order Management System", "Create", "IF-0002"),
        f("FLOW-0003", "Vehicle Configuration", "Internal", "APP-0003", "Product Catalog", "APP-0002", "Configure Price Quote", "Read", "IF-0004"),
        f("FLOW-0004", "Vehicle Order", "Internal", "APP-0005", "Order Management System", "APP-0024", "ERP Core", "Create", "IF-0001"),
        f("FLOW-0005", "ATP Quantity", "Internal", "APP-0030", "Inventory Management", "APP-0005", "Order Management System", "Read", "IF-0006"),
        f("FLOW-0006", "Billing Request", "Confidential", "APP-0005", "Order Management System", "APP-0006", "Billing System", "Create", "IF-0007"),
        f("FLOW-0007", "Customer Invoice", "Confidential", "APP-0006", "Billing System", "APP-0043", "Accounts Receivable", "Create", "IF-0008"),
        f("FLOW-0008", "Payment Instruction", "PCI", "APP-0006", "Billing System", "APP-0044", "Treasury", "Create", "IF-0016"),
        f("FLOW-0009", "Card Token", "PCI", "APP-0001", "Customer Portal", "APP-0006", "Billing System", "Create", "IF-0002"),
        f("FLOW-0010", "Production Confirmation", "Internal", "APP-0020", "Manufacturing Execution", "APP-0024", "ERP Core", "Update", "IF-0009"),
        f("FLOW-0011", "Engineering BOM", "Confidential", "APP-0021", "Product Lifecycle Management", "APP-0024", "ERP Core", "Upsert", "IF-0010"),
        f("FLOW-0012", "Pick Wave", "Internal", "APP-0005", "Order Management System", "APP-0031", "Warehouse Management", "Create", "IF-0012"),
        f("FLOW-0013", "Advanced Shipping Notice", "Internal", "APP-0033", "Supplier Portal", "APP-0031", "Warehouse Management", "Create", "IF-0013"),
        f("FLOW-0014", "Purchase Order", "Confidential", "APP-0034", "Procurement Suite", "APP-0024", "ERP Core", "Create", "IF-0014"),
        f("FLOW-0015", "Journal Entry", "Confidential", "APP-0024", "ERP Core", "APP-0041", "General Ledger", "Create", "IF-0015"),
        f("FLOW-0016", "Tax Quote", "Confidential", "APP-0045", "Tax Engine", "APP-0006", "Billing System", "Read", "IF-0017"),
        f("FLOW-0017", "Employee Record", "PII", "APP-0050", "HRIS Core", "APP-0051", "Payroll", "Upsert", "IF-0018"),
        f("FLOW-0018", "Bank Account", "PII", "APP-0050", "HRIS Core", "APP-0051", "Payroll", "Upsert", "IF-0018"),
        f("FLOW-0019", "Time Clocking", "PII", "APP-0053", "Time and Attendance", "APP-0051", "Payroll", "Create", "IF-0019"),
        f("FLOW-0020", "Party Golden Record", "PII", "APP-0062", "Customer Master Hub", "APP-0060", "CRM", "Upsert", "IF-0020"),
        f("FLOW-0021", "Marketing Lead", "PII", "APP-0061", "Marketing Automation", "APP-0060", "CRM", "Create", "IF-0021"),
        f("FLOW-0022", "Customer Snapshot", "PII", "APP-0062", "Customer Master Hub", "APP-0005", "Order Management System", "Upsert", "IF-0022"),
        f("FLOW-0023", "Order Fact", "Internal", "APP-0005", "Order Management System", "APP-0070", "Enterprise Data Warehouse", "Append", "IF-0023"),
        f("FLOW-0024", "Order Event", "Internal", "APP-0005", "Order Management System", "APP-0074", "Integration Streaming Bus", "Publish", "IF-0024"),
        f("FLOW-0025", "Auth Assertion", "Confidential", "APP-0090", "Identity Provider", "APP-0005", "Order Management System", "Create", "IF-0025"),
        f("FLOW-0026", "EDI 850 Purchase Order", "Confidential", "APP-0084", "Partner B2B Hub", "APP-0040", "Legacy EDI Gateway", "Create", "IF-0027"),
        f("FLOW-0027", "Dealer Vehicle Order", "Internal", "APP-0095", "Dealer DMS Adapter", "APP-0005", "Order Management System", "Create", "IF-0028"),
        f("FLOW-0028", "Line Telemetry", "Internal", "APP-0096", "Plant Historian", "APP-0071", "Data Lake", "Append", "IF-0029"),
        f("FLOW-0029", "Expense Report", "PII", "APP-0098", "Expense Management", "APP-0042", "Accounts Payable", "Create", "IF-0030"),
        f("FLOW-0030", "Supplier Contract", "Confidential", "APP-0099", "Contract Lifecycle", "APP-0034", "Procurement Suite", "Upsert", "IF-0031"),
        f("FLOW-0031", "Installed Base", "Internal", "APP-0021", "Product Lifecycle Management", "APP-0063", "Field Service", "Upsert", "IF-0032"),
        f("FLOW-0032", "Loyalty Profile", "PII", "APP-0062", "Customer Master Hub", "APP-0064", "Loyalty Platform", "Upsert", "IF-0033"),
        f("FLOW-0033", "Gateway Audit Log", "Confidential", "APP-0080", "API Gateway", "APP-0092", "SIEM", "Append", "IF-0034"),
        f("FLOW-0034", "Demand Plan", "Internal", "APP-0035", "Demand Planning", "APP-0025", "Production Scheduling", "Upsert", "IF-0036"),
        f("FLOW-0035", "Nonconformance", "Internal", "APP-0020", "Manufacturing Execution", "APP-0022", "Quality Management System", "Create", "IF-0037"),
        f("FLOW-0036", "Yard Visit", "Internal", "APP-0031", "Warehouse Management", "APP-0036", "Yard Management", "Create", "IF-0038"),
        f("FLOW-0037", "Vehicle Order", "Internal", "APP-0005", "Order Management System", "APP-0024", "ERP Core", "Update", "IF-9999"),
        f("FLOW-0038", "Contact Interaction", "PII", "APP-0060", "CRM", "APP-0065", "Contact Center", "Read", "IF-0040"),
        f("FLOW-0039", "Secret Reference", "Confidential", "APP-0093", "Secrets Vault", "APP-0081", "iPaaS", "Read", "IF-0035"),
        f("FLOW-0040", "Carbon Inventory", "Internal", "APP-0070", "Enterprise Data Warehouse", "APP-0097", "Carbon Accounting", "Read", ""),
    ]


def processes():
    def p(mid, pid, name, domain, app, aname, role, crit):
        return {
            "ProcessMappingID": mid,
            "BusinessProcessID": pid,
            "BusinessProcessName": name,
            "ProcessDomain": domain,
            "SupportingApplicationID": app,
            "SupportingApplicationName": aname,
            "RoleOfApplication": role,
            "ProcessCriticality": crit,
        }

    return [
        p("PM-0001", "BP-OTC-01", "Order Capture", "Order-to-Cash", "APP-0001", "Customer Portal", "Channel", "High"),
        p("PM-0002", "BP-OTC-01", "Order Capture", "Order-to-Cash", "APP-0002", "Configure Price Quote", "Configure", "High"),
        p("PM-0003", "BP-OTC-01", "Order Capture", "Order-to-Cash", "APP-0005", "Order Management System", "System of Record", "Mission Critical"),
        p("PM-0004", "BP-OTC-02", "Order Fulfilment", "Order-to-Cash", "APP-0005", "Order Management System", "Orchestrator", "Mission Critical"),
        p("PM-0005", "BP-OTC-02", "Order Fulfilment", "Order-to-Cash", "APP-0024", "ERP Core", "Manufacturing order", "Mission Critical"),
        p("PM-0006", "BP-OTC-02", "Order Fulfilment", "Order-to-Cash", "APP-0030", "Inventory Management", "Availability", "High"),
        p("PM-0007", "BP-OTC-02", "Order Fulfilment", "Order-to-Cash", "APP-0031", "Warehouse Management", "Pick/ship", "High"),
        p("PM-0008", "BP-OTC-03", "Bill to Cash", "Order-to-Cash", "APP-0006", "Billing System", "System of Record", "High"),
        p("PM-0009", "BP-OTC-03", "Bill to Cash", "Order-to-Cash", "APP-0007", "Invoice Management", "Presentment", "Medium"),
        p("PM-0010", "BP-OTC-03", "Bill to Cash", "Order-to-Cash", "APP-0043", "Accounts Receivable", "Open items", "High"),
        p("PM-0011", "BP-MFG-01", "Production Execution", "Manufacturing", "APP-0020", "Manufacturing Execution", "System of Record", "Mission Critical"),
        p("PM-0012", "BP-MFG-01", "Production Execution", "Manufacturing", "APP-0023", "Shop Floor Control", "Execution", "High"),
        p("PM-0013", "BP-MFG-01", "Production Execution", "Manufacturing", "APP-0024", "ERP Core", "Order", "Mission Critical"),
        p("PM-0014", "BP-MFG-02", "Quality Containment", "Manufacturing", "APP-0022", "Quality Management System", "System of Record", "High"),
        p("PM-0015", "BP-MFG-02", "Quality Containment", "Manufacturing", "APP-0020", "Manufacturing Execution", "Containment", "High"),
        p("PM-0016", "BP-SCM-01", "Procure to Pay", "Supply Chain", "APP-0034", "Procurement Suite", "System of Record", "High"),
        p("PM-0017", "BP-SCM-01", "Procure to Pay", "Supply Chain", "APP-0042", "Accounts Payable", "Settlement", "High"),
        p("PM-0018", "BP-SCM-01", "Procure to Pay", "Supply Chain", "APP-0024", "ERP Core", "Goods receipt", "High"),
        p("PM-0019", "BP-SCM-02", "Inbound Logistics", "Supply Chain", "APP-0031", "Warehouse Management", "Receiving", "High"),
        p("PM-0020", "BP-SCM-02", "Inbound Logistics", "Supply Chain", "APP-0033", "Supplier Portal", "ASN", "Medium"),
        p("PM-0021", "BP-FIN-01", "Record to Report", "Finance", "APP-0041", "General Ledger", "System of Record", "Mission Critical"),
        p("PM-0022", "BP-FIN-01", "Record to Report", "Finance", "APP-0024", "ERP Core", "Subledger", "Mission Critical"),
        p("PM-0023", "BP-FIN-02", "Treasury Payments", "Finance", "APP-0044", "Treasury", "System of Record", "High"),
        p("PM-0024", "BP-FIN-02", "Treasury Payments", "Finance", "APP-0006", "Billing System", "Source", "High"),
        p("PM-0025", "BP-HR-01", "Hire to Retire", "HR", "APP-0050", "HRIS Core", "System of Record", "High"),
        p("PM-0026", "BP-HR-01", "Hire to Retire", "HR", "APP-0051", "Payroll", "Pay", "Mission Critical"),
        p("PM-0027", "BP-HR-01", "Hire to Retire", "HR", "APP-0052", "Recruiting", "Acquire", "Medium"),
        p("PM-0028", "BP-CX-01", "Lead to Cash", "Customer", "APP-0060", "CRM", "Lead/account", "High"),
        p("PM-0029", "BP-CX-01", "Lead to Cash", "Customer", "APP-0005", "Order Management System", "Order", "Mission Critical"),
        p("PM-0030", "BP-CX-01", "Lead to Cash", "Customer", "APP-0006", "Billing System", "Invoice", "High"),
        p("PM-0031", "BP-CX-02", "Customer Service", "Customer", "APP-0065", "Contact Center", "Channel", "Medium"),
        p("PM-0032", "BP-CX-02", "Customer Service", "Customer", "APP-0060", "CRM", "Case", "High"),
        p("PM-0033", "BP-DAT-01", "Enterprise Reporting", "Data Platform", "APP-0070", "Enterprise Data Warehouse", "System of Record", "High"),
        p("PM-0034", "BP-DAT-01", "Enterprise Reporting", "Data Platform", "APP-0073", "Analytics Workbench", "Consume", "Medium"),
        p("PM-0035", "BP-SEC-01", "Workforce Access", "Identity/Security", "APP-0090", "Identity Provider", "System of Record", "Mission Critical"),
        p("PM-0036", "BP-SEC-01", "Workforce Access", "Identity/Security", "APP-0080", "API Gateway", "Enforce", "High"),
        p("PM-0037", "BP-OTC-02", "Order Fulfilment", "Order-to-Cash", "APP-0040", "Legacy EDI Gateway", "Partner intake", "High"),
        p("PM-0038", "BP-OTC-01", "Order Capture", "Order-to-Cash", "APP-8824", "ERP Core", "Pricing fallback", "High"),
        p("PM-0039", "BP-CX-03", "Loyalty Earn and Burn", "Customer", "APP-0064", "Loyalty Platform", "System of Record", "Low"),
        p("PM-0040", "BP-MFG-03", "Maintenance Work", "Manufacturing", "APP-0026", "Maintenance Management", "System of Record", "Medium"),
    ]


def ownership():
    def o(oid, aid, aname, owner, eid, custodian, biz, support, dept):
        return {
            "OwnershipID": oid,
            "ApplicationID": aid,
            "ApplicationName": aname,
            "ApplicationOwner": owner,
            "OwnerEmployeeID": eid,
            "SystemCustodian": custodian,
            "BusinessOwner": biz,
            "SupportGroup": support,
            "Department": dept,
        }

    # Intentionally omit APP-0061, APP-0054, APP-0036, APP-0097.
    # Blank owner on APP-0082.
    rows = [
        o("OWN-0001", "APP-0001", "Customer Portal", "Priya Raman", "E-1041", "Digital Ops", "VP Retail", "L2 Digital", "Order-to-Cash"),
        o("OWN-0002", "APP-0002", "Configure Price Quote", "Priya Raman", "E-1041", "Digital Ops", "VP Retail", "L2 Digital", "Order-to-Cash"),
        o("OWN-0003", "APP-0003", "Product Catalog", "James Okoye", "E-1102", "Product IT", "Head of Product", "L2 Product", "Order-to-Cash"),
        o("OWN-0004", "APP-0004", "Pricing Engine", "James Okoye", "E-1102", "Product IT", "Head of Product", "L2 Product", "Order-to-Cash"),
        o("OWN-0005", "APP-0005", "Order Management System", "Elena Voss", "E-1007", "OMS Platform", "VP Order Management", "L2 OMS", "Order-to-Cash"),
        o("OWN-0006", "APP-0006", "Billing System", "Chen Wei", "E-1220", "Finance IT", "CFO Office", "L2 Finance", "Finance"),
        o("OWN-0007", "APP-0007", "Invoice Management", "Chen Wei", "E-1220", "Finance IT", "CFO Office", "L2 Finance", "Finance"),
        o("OWN-0008", "APP-0008", "Collections Workbench", "Amira Haddad", "E-1221", "Finance IT", "CFO Office", "L2 Finance", "Finance"),
        o("OWN-0009", "APP-0020", "Manufacturing Execution", "Lars Holm", "E-2001", "Plant IT", "VP Manufacturing", "L2 MES", "Manufacturing"),
        o("OWN-0010", "APP-0021", "Product Lifecycle Management", "Sofia Berg", "E-2008", "Engineering IT", "CTO Office", "L2 PLM", "Engineering"),
        o("OWN-0011", "APP-0022", "Quality Management System", "Nina Park", "E-2014", "Quality IT", "VP Quality", "L2 QMS", "Manufacturing"),
        o("OWN-0012", "APP-0023", "Shop Floor Control", "Lars Holm", "E-2001", "Plant IT", "VP Manufacturing", "L2 MES", "Manufacturing"),
        o("OWN-0013", "APP-0024", "ERP Core", "Martin Shaw", "E-3001", "ERP Centre", "CFO / COO", "L1 ERP", "Enterprise"),
        o("OWN-0014", "APP-0025", "Production Scheduling", "Ines Duarte", "E-2003", "Plant IT", "VP Manufacturing", "L2 Planning", "Manufacturing"),
        o("OWN-0015", "APP-0026", "Maintenance Management", "Hugo Klein", "E-2050", "Plant IT", "VP Manufacturing", "L2 Maint", "Manufacturing"),
        o("OWN-0016", "APP-0030", "Inventory Management", "Rita Gomez", "E-3104", "Supply IT", "VP Supply Chain", "L2 Inventory", "Supply Chain"),
        o("OWN-0017", "APP-0031", "Warehouse Management", "Rita Gomez", "E-3104", "Supply IT", "VP Supply Chain", "L2 WMS", "Supply Chain"),
        o("OWN-0018", "APP-0032", "Transportation Management", "Owen Blake", "E-3110", "Supply IT", "VP Supply Chain", "L2 TMS", "Supply Chain"),
        o("OWN-0019", "APP-0033", "Supplier Portal", "Mei Lin", "E-3122", "Procurement IT", "CPO", "L2 Supplier", "Supply Chain"),
        o("OWN-0020", "APP-0034", "Procurement Suite", "Mei Lin", "E-3120", "Procurement IT", "CPO", "L2 P2P", "Supply Chain"),
        o("OWN-0021", "APP-0035", "Demand Planning", "Chris Adeyemi", "E-3130", "Supply IT", "VP Supply Chain", "L2 Planning", "Supply Chain"),
        o("OWN-0022", "APP-0040", "Legacy EDI Gateway", "Retired Run Team", "E-4099", "Integration", "VP Integration", "L3 EDI", "Integration"),
        o("OWN-0023", "APP-0041", "General Ledger", "Helen Cho", "E-4001", "Finance IT", "Controller", "L1 Finance", "Finance"),
        o("OWN-0024", "APP-0042", "Accounts Payable", "Helen Cho", "E-4002", "Finance IT", "Controller", "L2 AP", "Finance"),
        o("OWN-0025", "APP-0043", "Accounts Receivable", "Helen Cho", "E-4003", "Finance IT", "Controller", "L2 AR", "Finance"),
        o("OWN-0026", "APP-0044", "Treasury", "Noah Stein", "E-4010", "Finance IT", "Treasurer", "L2 Treasury", "Finance"),
        o("OWN-0027", "APP-0045", "Tax Engine", "Ivy Grant", "E-4011", "Finance IT", "Tax Director", "L2 Tax", "Finance"),
        o("OWN-0028", "APP-0050", "HRIS Core", "Paula Nunez", "E-5001", "HR IT", "CHRO", "L2 HR", "HR"),
        o("OWN-0029", "APP-0051", "Payroll", "Paula Nunez", "E-5002", "HR IT", "CHRO", "L1 Payroll", "HR"),
        o("OWN-0030", "APP-0052", "Recruiting", "Tomasi Ratu", "E-5003", "HR IT", "CHRO", "L2 TA", "HR"),
        o("OWN-0031", "APP-0053", "Time and Attendance", "Paula Nunez", "E-5004", "HR IT", "CHRO", "L2 HR", "HR"),
        o("OWN-0032", "APP-0060", "CRM", "Aisha Rahman", "E-6001", "CX IT", "CCO", "L2 CRM", "Customer"),
        o("OWN-0033", "APP-0062", "Customer Master Hub", "Aisha Rahman", "E-6002", "CX IT", "CCO", "L2 MDM", "Customer"),
        o("OWN-0034", "APP-0063", "Field Service", "Benito Cruz", "E-6008", "CX IT", "VP Service", "L2 FS", "Customer"),
        o("OWN-0035", "APP-0064", "Loyalty Platform", "Keiko Mori", "E-6011", "CX IT", "CMO", "L3 Loyalty", "Customer"),
        o("OWN-0036", "APP-0065", "Contact Center", "Leila Haddad", "E-6012", "CX IT", "CCO", "L2 CC", "Customer"),
        o("OWN-0037", "APP-0070", "Enterprise Data Warehouse", "Samir Patel", "E-7001", "Data Platform", "CDO", "L1 EDW", "Data"),
        o("OWN-0038", "APP-0071", "Data Lake", "Samir Patel", "E-7001", "Data Platform", "CDO", "L1 Lake", "Data"),
        o("OWN-0039", "APP-0072", "Master Data Management", "Grace Cho", "E-7004", "Data Platform", "CDO", "L2 MDM", "Data"),
        o("OWN-0040", "APP-0073", "Analytics Workbench", "Ibrahim Sule", "E-7008", "Data Platform", "CDO", "L3 BI", "Data"),
        o("OWN-0041", "APP-0074", "Integration Streaming Bus", "Nora Quint", "E-7010", "Integration", "VP Integration", "L1 Events", "Integration"),
        o("OWN-0042", "APP-0075", "Data Catalog", "Ibrahim Sule", "E-7009", "Data Platform", "CDO", "L3 Catalog", "Data"),
        o("OWN-0043", "APP-0080", "API Gateway", "Nora Quint", "E-8001", "Integration", "VP Integration", "L1 API", "Integration"),
        o("OWN-0044", "APP-0081", "iPaaS", "Nora Quint", "E-8001", "Integration", "VP Integration", "L2 iPaaS", "Integration"),
        o("OWN-0045", "APP-0082", "File Transfer Gateway", "", "", "Integration", "VP Integration", "L3 MFT", "Integration"),
        o("OWN-0046", "APP-0083", "Event Broker", "Nora Quint", "E-7010", "Integration", "VP Integration", "L1 Events", "Integration"),
        o("OWN-0047", "APP-0084", "Partner B2B Hub", "Diego Alves", "E-8006", "Integration", "VP Integration", "L2 B2B", "Integration"),
        o("OWN-0048", "APP-0090", "Identity Provider", "Hannah Cole", "E-9001", "Security", "CISO", "L1 IAM", "Security"),
        o("OWN-0049", "APP-0091", "Privileged Access", "Hannah Cole", "E-9002", "Security", "CISO", "L1 IAM", "Security"),
        o("OWN-0050", "APP-0092", "SIEM", "Yusef Farouk", "E-9003", "Security", "CISO", "L1 SOC", "Security"),
        o("OWN-0051", "APP-0093", "Secrets Vault", "Hannah Cole", "E-9002", "Security", "CISO", "L1 IAM", "Security"),
        o("OWN-0052", "APP-0094", "GRC Platform", "Clara Jensen", "E-9008", "Security", "CISO", "L2 GRC", "Security"),
        o("OWN-0053", "APP-0095", "Dealer DMS Adapter", "Elena Voss", "E-1044", "OMS Platform", "VP Order Management", "L2 OMS", "Order-to-Cash"),
        o("OWN-0054", "APP-0096", "Plant Historian", "Lars Holm", "E-2060", "Plant IT", "VP Manufacturing", "L3 OT", "Manufacturing"),
        o("OWN-0055", "APP-0098", "Expense Management", "Helen Cho", "E-4022", "Finance IT", "Controller", "L3 T&E", "Finance"),
        o("OWN-0056", "APP-0099", "Contract Lifecycle", "Mei Lin", "E-3128", "Procurement IT", "CPO", "L3 CLM", "Supply Chain"),
    ]
    return rows


def gaps():
    def g(gid, gtype, etype, eid, rel, desc, sev):
        return {
            "GapID": gid,
            "GapType": gtype,
            "EntityType": etype,
            "EntityID": eid,
            "RelatedApplicationID": rel,
            "Description": desc,
            "Severity": sev,
        }

    return [
        g("GAP-0001", "MissingOwner", "Application", "APP-0061", "APP-0061", "Marketing Automation has no ApplicationOwnership row.", "High"),
        g("GAP-0002", "StaleInterface", "Interface", "IF-0027", "APP-0040", "Partner EDI path still marked in use on a Retired gateway.", "High"),
        g("GAP-0003", "UnclearSystemOfRecord", "InformationObject", "FLOW-0004", "APP-0005", "Vehicle Order is created in OMS and ERP; SoR not agreed.", "Medium"),
        g("GAP-0004", "DuplicateCustomerKey", "InformationObject", "FLOW-0020", "APP-0062", "CRM still assigns local party IDs alongside golden records.", "Medium"),
        g("GAP-0005", "MissingInterface", "Application", "APP-0054", "APP-0054", "Learning Management has no recorded interfaces or relationships.", "Low"),
        g("GAP-0006", "PCIScopeUndocumented", "InformationObject", "FLOW-0008", "APP-0044", "Payment Instruction classified PCI but no tokenisation control recorded.", "High"),
        g("GAP-0007", "LifecycleMismatch", "Application", "APP-0052", "APP-0052", "Recruiting start date is after the recorded end date.", "Medium"),
        g("GAP-0008", "BrokenReference", "Relationship", "REL-0063", "APP-0005", "OMS dependency points at APP-8824 which is not in Applications.", "High"),
        g("GAP-0009", "OrphanInterface", "Interface", "IF-0088", "APP-0003", "Catalog ping interface has no information flow rows.", "Low"),
        g("GAP-0010", "PIIDownstream", "InformationObject", "FLOW-0021", "APP-0061", "Marketing lead PII lands in CRM from an application past its lifecycle end.", "High"),
    ]


SHEETS = [
    ("Applications", applications),
    ("Relationships", relationships),
    ("Interfaces", interfaces),
    ("InformationObjects", information_objects),
    ("BusinessProcesses", processes),
    ("ApplicationOwnership", ownership),
    ("KnownDataQualityGaps", gaps),
]


def write_workbook(path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    default = wb.active
    wb.remove(default)
    for title, factory in SHEETS:
        rows = factory()
        ws = wb.create_sheet(title)
        headers = list(rows[0].keys())
        ws.append(headers)
        for cell in ws[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center")
        for row in rows:
            values = []
            for h in headers:
                v = row[h]
                if isinstance(v, date):
                    v = v.isoformat()
                values.append(v)
            ws.append(values)
            for cell in ws[ws.max_row]:
                cell.font = CELL_FONT
        for idx, header in enumerate(headers, start=1):
            width = max(len(header) + 2, 18)
            for cell in ws[get_column_letter(idx)]:
                if cell.value:
                    width = max(width, min(len(str(cell.value)) + 2, 42))
            ws.column_dimensions[get_column_letter(idx)].width = width
        ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes = "A2"
    wb.save(path)
    return path


def main(argv=None):
    dest = Path(argv[1]) if argv and len(argv) > 1 else OUT
    write_workbook(dest)
    print(f"Wrote {dest}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
